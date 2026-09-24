"""PSX portfolio tracker. Serves the sheet and refreshes live market data."""

from __future__ import annotations

import calendar
import json
import os
import re
import threading
import urllib.error
import urllib.request
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
INPUTS = DATA / "inputs.json"
CACHE = DATA / "cache.json"
STATIC = ROOT / "static"
PKT = ZoneInfo("Asia/Karachi")
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/json",
}
FREQ_MAP = {
    "quarterly": ("Quarterly", "Q"),
    "semi-annual": ("Semi Annual", "SA"),
    "semiannual": ("Semi Annual", "SA"),
    "semi annual": ("Semi Annual", "SA"),
    "annual": ("Annual", "Ann."),
}

_lock = threading.Lock()
_cache = {"market": {}, "stats": {}, "eod": {}, "div": {}}
_snapshot = None
_refreshing = False
_refresh_error = None


def now_pkt():
    return datetime.now(PKT)


def stamp():
    return now_pkt().strftime("%d %b %Y, %I:%M %p PKT")


def market_status():
    current = now_pkt()
    opens = current.replace(hour=9, minute=30, second=0, microsecond=0)
    closes = current.replace(hour=15, minute=30, second=0, microsecond=0)
    open_now = current.weekday() < 5 and opens <= current <= closes
    if open_now:
        return "Market is open, will close at 3:30 PM"
    if current.weekday() < 4 or (current.weekday() == 4 and current < opens):
        when = "today" if current < opens else "tomorrow"
        return f"Market is closed, will open {when} at 9:30 AM"
    return "Market is closed, will open Monday at 9:30 AM"


def load_inputs():
    return json.loads(INPUTS.read_text(encoding="utf-8"))["stocks"]


def save_inputs(stocks):
    symbols = [s["symbol"] for s in stocks]
    if len(symbols) != len(set(symbols)):
        raise ValueError("Duplicate symbols are not allowed.")
    INPUTS.write_text(
        json.dumps({"stocks": stocks}, indent=2), encoding="utf-8"
    )


def load_cache():
    global _cache
    if CACHE.exists():
        _cache = json.loads(CACHE.read_text(encoding="utf-8"))
        for key in ("market", "stats", "eod", "div"):
            _cache.setdefault(key, {})


def save_cache():
    CACHE.write_text(json.dumps(_cache), encoding="utf-8")


def fetch(url):
    req = urllib.request.Request(url, headers=HEADERS)
    with urllib.request.urlopen(req, timeout=30) as response:
        return response.read().decode("utf-8", "replace")


def parse_pct(text):
    if not text:
        return None
    match = re.search(r"(-?\d+(?:\.\d+)?)\s*%", text.replace(",", ""))
    if not match:
        return None
    return float(match.group(1)) / 100


def parse_market(html):
    found = {}
    for row in re.findall(r"<tr>(.*?)</tr>", html, re.S):
        symbol = re.search(r'data-search="([A-Z0-9]+)"', row)
        orders = re.findall(r'data-order="([^"]*)"', row)
        if not symbol or len(orders) < 5:
            continue
        try:
            found[symbol.group(1)] = {
                "ldcp": float(orders[1]),
                "price": float(orders[5]),
            }
        except (ValueError, IndexError):
            continue
    return found


def parse_company(html):
    week = re.search(
        r"52-WEEK RANGE.*?data-low=\"([^\"]+)\"\s+data-high=\"([^\"]+)\"",
        html,
        re.S,
    )
    ytd = re.search(r"YTD Change.*?<div class=\"stats_value[^\"]*\">(.*?)</div>", html, re.S)
    year = re.search(
        r"1-Year Change.*?<div class=\"stats_value[^\"]*\">(.*?)</div>", html, re.S
    )
    high = low = None
    if week:
        low = float(week.group(1))
        high = float(week.group(2))
    ytd_pct = parse_pct(re.sub(r"<[^>]+>", "", ytd.group(1))) if ytd else None
    year_pct = parse_pct(re.sub(r"<[^>]+>", "", year.group(1))) if year else None
    return {"high52": high, "low52": low, "ytd": ytd_pct, "year": year_pct}


def parse_eod(payload):
    data = json.loads(payload).get("data") or []
    series = []
    for row in data:
        if len(row) >= 2:
            series.append([int(row[0]), float(row[1])])
    return series


def infer_frequency(history):
    if len(history) < 2:
        return "Annual", "Ann."
    dates = sorted(datetime.fromisoformat(item["date"]) for item in history[:4])
    gaps = [(dates[index + 1] - dates[index]).days for index in range(len(dates) - 1)]
    median = sorted(gaps)[len(gaps) // 2]
    if median < 140:
        return "Quarterly", "Q"
    if median < 270:
        return "Semi Annual", "SA"
    return "Annual", "Ann."


def parse_dividend(html):
    frequency = re.search(r'frequency:"([^"]*)"', html)
    yield_text = re.search(r'infoTable:\{yield:"([^"]*)"', html)
    history = []
    for date, amount in re.findall(r'\{dt:"(\d{4}-\d{2}-\d{2})",amt:"([^"]+)"', html):
        number = re.search(r"(-?\d+(?:\.\d+)?)", amount.replace(",", ""))
        if number:
            history.append({"date": date, "amount": float(number.group(1))})
    label = frequency.group(1).strip() if frequency else ""
    mapped = FREQ_MAP.get(label.lower())
    if not mapped and label.lower() in {"", "n/a", "none"}:
        mapped = infer_frequency(history) if history else ("None", None)
    raw_yield = yield_text.group(1).strip() if yield_text else ""
    if raw_yield.lower() == "n/a":
        div_yield = 0.0
    else:
        div_yield = parse_pct(raw_yield) if raw_yield else None
    return {
        "frequency": mapped[0] if mapped else (label or None),
        "marker": mapped[1] if mapped else None,
        "yield": div_yield,
        "history": history,
    }


def shift_months(moment, months):
    month_index = moment.month - 1 + months
    year = moment.year + month_index // 12
    month = month_index % 12 + 1
    day = min(moment.day, calendar.monthrange(year, month)[1])
    return moment.replace(year=year, month=month, day=day)


def close_on_or_before(series, target):
    target_ts = int(target.timestamp())
    chosen = None
    for ts, close in series:
        if ts <= target_ts:
            chosen = (ts, close)
            break
    return chosen


def payout_cells(entry):
    if not entry or entry.get("missing"):
        return "None", ["None", "None", "None", "None"]
    history = entry.get("history") or []
    if not history:
        return "None", ["None", "None", "None", "None"]
    today = now_pkt().date().isoformat()
    past = [item for item in history if item["date"] <= today]
    chosen = past[:4] if past else history[:4]
    marker = entry.get("marker")
    frequency = entry.get("frequency") or "Quarterly"
    cells = []
    for item in chosen:
        suffix = f" ({marker})" if marker else ""
        cells.append(f"Rs. {item['amount']:.2f}{suffix}")
    while len(cells) < 4:
        cells.append("N/A")
    return frequency, cells


def gain_from(price, base):
    if price is None or base is None or base == 0:
        return None
    return price / base - 1


def refresh_symbol(symbol, market):
    result = {"symbol": symbol}
    try:
        result["stats"] = parse_company(fetch(f"https://dps.psx.com.pk/company/{symbol}"))
        result["stats_ok"] = True
    except (urllib.error.URLError, TimeoutError, ValueError):
        result["stats_ok"] = False
    try:
        result["eod"] = parse_eod(fetch(f"https://dps.psx.com.pk/timeseries/eod/{symbol}"))
        result["eod_ok"] = bool(result["eod"])
    except (urllib.error.URLError, TimeoutError, ValueError, json.JSONDecodeError):
        result["eod_ok"] = False
    try:
        result["div"] = parse_dividend(
            fetch(f"https://stockanalysis.com/quote/psx/{symbol}/dividend/")
        )
        result["div_ok"] = True
    except urllib.error.HTTPError as error:
        if error.code == 404:
            result["div"] = {"frequency": "None", "marker": None, "yield": 0.0, "history": []}
            result["div_ok"] = True
        else:
            result["div_ok"] = False
    except (urllib.error.URLError, TimeoutError, ValueError):
        result["div_ok"] = False
    result["market_ok"] = symbol in market
    return result


def keep(previous, fresh, ok):
    if ok and fresh is not None:
        return fresh, False
    if previous is not None:
        return previous, True
    return None, True


def build_snapshot():
    stocks = load_inputs()
    symbols = [row["symbol"] for row in stocks]
    stale = False
    try:
        market = parse_market(fetch("https://dps.psx.com.pk/market-watch"))
        market_ok = True
    except (urllib.error.URLError, TimeoutError):
        market = {}
        market_ok = False
        stale = True

    fetched = {}
    with ThreadPoolExecutor(max_workers=12) as pool:
        futures = [pool.submit(refresh_symbol, symbol, market) for symbol in symbols]
        for future in as_completed(futures):
            item = future.result()
            fetched[item["symbol"]] = item

    current = now_pkt()
    three_months = shift_months(current, -3)
    year_ago = current - timedelta(weeks=52)
    year_start = datetime(current.year, 1, 1, tzinfo=PKT)
    rows = []
    dividend_log = {}

    for stock in stocks:
        symbol = stock["symbol"]
        item = fetched.get(symbol, {})
        previous_market = _cache["market"].get(symbol, {})
        previous_stats = _cache["stats"].get(symbol, {})
        previous_eod = _cache["eod"].get(symbol, {})
        previous_div = _cache["div"].get(symbol, {})

        quote = market.get(symbol, {})
        ldcp, ldcp_stale = keep(previous_market.get("ldcp"), quote.get("ldcp"), market_ok and symbol in market)
        price, price_stale = keep(previous_market.get("price"), quote.get("price"), market_ok and symbol in market)
        if ldcp is not None:
            _cache["market"][symbol] = {"ldcp": ldcp, "price": price}

        stats = item.get("stats") or {}
        high, high_stale = keep(previous_stats.get("high52"), stats.get("high52"), item.get("stats_ok"))
        low, low_stale = keep(previous_stats.get("low52"), stats.get("low52"), item.get("stats_ok"))
        ytd, ytd_stale = keep(previous_stats.get("ytd"), stats.get("ytd"), item.get("stats_ok"))
        year, year_stale = keep(previous_stats.get("year"), stats.get("year"), item.get("stats_ok"))

        series = item.get("eod") if item.get("eod_ok") else previous_eod.get("series")
        eod_stale = not item.get("eod_ok")
        base_3m = previous_eod.get("base3m")
        base_1y = previous_eod.get("base1y")
        base_ytd = previous_eod.get("baseYtd")
        if series:
            found_3m = close_on_or_before(series, three_months)
            found_1y = close_on_or_before(series, year_ago)
            found_ytd = close_on_or_before(series, year_start)
            base_3m = found_3m[1] if found_3m else None
            base_1y = found_1y[1] if found_1y else None
            base_ytd = found_ytd[1] if found_ytd else None
        _cache["eod"][symbol] = {"base3m": base_3m, "base1y": base_1y, "baseYtd": base_ytd}
        gain_3m = gain_from(price, base_3m)
        if gain_3m is None and eod_stale:
            gain_3m, gain_3m_stale = previous_stats.get("gain3m"), True
        else:
            gain_3m_stale = eod_stale or price_stale
        if ytd is None:
            ytd, ytd_stale = keep(previous_stats.get("ytd"), gain_from(price, base_ytd), not eod_stale)
        if year is None:
            year, year_stale = keep(previous_stats.get("year"), gain_from(price, base_1y), not eod_stale)

        _cache["stats"][symbol] = {
            "high52": high,
            "low52": low,
            "ytd": ytd,
            "year": year,
            "gain3m": gain_3m,
        }

        if item.get("div_ok"):
            dividend = item.get("div")
            _cache["div"][symbol] = dividend
            div_stale = False
        else:
            dividend = previous_div or None
            div_stale = True
        frequency, payouts = payout_cells(dividend)
        div_yield = (dividend or {}).get("yield")
        if div_yield is None and div_stale:
            div_yield = previous_div.get("yield")
        dividend_log[symbol] = (dividend or {}).get("history") or []

        flags = []
        if price is None or price <= 0:
            flags.append("Price missing")
        elif ldcp and ldcp > 0 and abs(price - ldcp) / ldcp > 0.20:
            flags.append("Price moved more than 20% from LDCP")

        field_stale = {
            "ldcp": ldcp_stale,
            "price": price_stale,
            "yield": div_stale or div_yield is None,
            "frequency": div_stale,
            "payouts": div_stale,
            "gain3m": gain_3m_stale or gain_3m is None,
            "ytd": ytd_stale or ytd is None,
            "year": year_stale or year is None,
            "high52": high_stale or high is None,
            "low52": low_stale or low is None,
        }
        if any(field_stale.values()):
            stale = True

        rows.append(
            {
                **stock,
                "ldcp": ldcp,
                "price": price,
                "yield": div_yield,
                "frequency": frequency,
                "payouts": payouts,
                "gain3m": gain_3m,
                "ytd": ytd,
                "year": year,
                "high52": high,
                "low52": low,
                "base3m": base_3m,
                "base1y": base_1y,
                "baseYtd": base_ytd,
                "flags": flags,
                "stale": field_stale,
                "history": dividend_log[symbol],
            }
        )

    save_cache()
    status = market_status()
    return {
        "updated": stamp(),
        "market": status,
        "marketOpen": status.startswith("Market is open"),
        "stale": stale,
        "rows": rows,
    }


def refresh_in_background():
    global _refreshing

    with _lock:
        if _refreshing:
            return
        _refreshing = True

    def run():
        global _snapshot, _refreshing, _refresh_error
        try:
            snapshot = build_snapshot()
            with _lock:
                _snapshot = snapshot
                _refresh_error = None
        except Exception as error:  # noqa: BLE001 - keep the site up if a fetch fails
            with _lock:
                _refresh_error = str(error)
        finally:
            with _lock:
                _refreshing = False

    threading.Thread(target=run, daemon=True).start()


PCT_COLS = {4, 10, 15, 17, 23, 24, 25}
MONEY_COLS = {5, 6, 7, 8, 9, 12, 13, 14, 16, 26, 27}
GAIN_COLS = {15, 17, 23, 24, 25}


def indian_text(value, digits=2):
    if value is None:
        return "—"
    negative = value < 0
    number = f"{abs(value):.{digits}f}"
    whole, fraction = number.split(".") if "." in number else (number, "")
    head, tail = whole[:-3], whole[-3:]
    if head:
        head = re.sub(r"\B(?=(\d{2})+(?!\d))", ",", head) + ","
    text = head + tail + (("." + fraction) if digits else "")
    return f"-{text}" if negative else text


def display_cell(value, index):
    if isinstance(value, str) or value is None:
        return "—" if value is None else str(value)
    if index in PCT_COLS:
        return f"{value * 100:.2f}%"
    if index in MONEY_COLS:
        return indian_text(value, 2)
    if index == 11:
        return indian_text(value, 0)
    if isinstance(value, float):
        return indian_text(value, 2)
    return str(value)


def build_xlsx(payload):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = "Portfolio"
    headers = payload["headers"]
    sheet["A1"] = "PSX Portfolio Tracker"
    sheet["A2"] = f"Total Portfolio Amount (Rs.): {indian_text(payload.get('amount') or 0, 2)}"
    sheet["A3"] = f"Last updated {payload.get('updated', '')}  |  Market {payload.get('market', '')}"
    sheet.append([])
    header_row = 5
    sheet.append(headers)
    for row in payload["rows"]:
        sheet.append(row)
    if payload.get("total"):
        sheet.append(payload["total"])
    header_fill = PatternFill("solid", fgColor="21493C")
    total_fill = PatternFill("solid", fgColor="EFE6D2")
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="F6F1E6")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
    last = sheet.max_row
    for col in range(1, len(headers) + 1):
        for row in range(header_row + 1, last + 1):
            cell = sheet.cell(row, col)
            if col - 1 in PCT_COLS and isinstance(cell.value, (int, float)):
                cell.number_format = "0.00%"
            elif col - 1 in MONEY_COLS and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##,##0.00"
            elif col == 12 and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##,##0"
            if col - 1 in GAIN_COLS and isinstance(cell.value, (int, float)):
                cell.font = Font(color="0C7A3E" if cell.value > 0 else "B42318" if cell.value < 0 else "1B241C")
        sheet.column_dimensions[get_column_letter(col)].width = 16
    for col, cell in enumerate(sheet[last], start=1):
        value = cell.value
        color = "1B241C"
        if col - 1 in GAIN_COLS and isinstance(value, (int, float)):
            color = "0C7A3E" if value > 0 else "B42318" if value < 0 else "1B241C"
        cell.font = Font(bold=True, color=color)
        cell.fill = total_fill
    sheet.freeze_panes = "C6"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last}"
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def build_pdf(payload):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer, pagesize=landscape(A3), leftMargin=18, rightMargin=18, topMargin=18, bottomMargin=18
    )
    title = ParagraphStyle("title", fontName="Times-Bold", fontSize=14, textColor=colors.HexColor("#16352C"))
    meta = ParagraphStyle("meta", fontName="Times-Roman", fontSize=9, textColor=colors.HexColor("#5C675C"))
    headers = payload["headers"]
    data = [headers]
    for row in payload["rows"]:
        data.append([display_cell(value, index) for index, value in enumerate(row)])
    if payload.get("total"):
        data.append([display_cell(value, index) for index, value in enumerate(payload["total"])])
    table = Table(data, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#21493C")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#F6F1E6")),
        ("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
        ("FONTNAME", (0, 1), (-1, -2), "Times-Roman"),
        ("FONTSIZE", (0, 0), (-1, -1), 6),
        ("ALIGN", (4, 1), (-1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#D9D1C2")),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#EFE6D2")),
        ("FONTNAME", (0, -1), (-1, -1), "Times-Bold"),
    ]
    for row_index, row in enumerate(payload["rows"], start=1):
        for col in GAIN_COLS:
            value = row[col] if col < len(row) else None
            if isinstance(value, (int, float)) and value != 0:
                style.append(("TEXTCOLOR", (col, row_index), (col, row_index), colors.HexColor("#0C7A3E" if value > 0 else "#B42318")))
    table.setStyle(TableStyle(style))
    document.build([
        Paragraph("PSX Portfolio Tracker", title),
        Spacer(1, 6),
        Paragraph(
            f"Total Portfolio Amount (Rs.): {indian_text(payload.get('amount') or 0, 2)} &nbsp;&nbsp; Last updated {payload.get('updated', '')} &nbsp;&nbsp; Market {payload.get('market', '')}",
            meta,
        ),
        Spacer(1, 8),
        table,
    ])
    return buffer.getvalue()


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(STATIC), **kwargs)

    def end_headers(self):
        self.send_header("Cache-Control", "no-store")
        super().end_headers()

    def do_GET(self):
        if self.path.startswith("/api/portfolio"):
            self.send_portfolio()
            return
        if self.path.startswith("/api/inputs"):
            self.send_json({"stocks": load_inputs()})
            return
        super().do_GET()

    def do_POST(self):
        length = int(self.headers.get("Content-Length", "0"))
        body = json.loads(self.rfile.read(length) or b"{}")
        if self.path.startswith("/api/inputs"):
            try:
                save_inputs(body["stocks"])
            except (KeyError, TypeError, ValueError) as error:
                self.send_json({"error": str(error)}, 400)
                return
            self.send_json({"ok": True})
            return
        if self.path.startswith("/api/export.xlsx"):
            self.send_bytes(
                build_xlsx(body),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "PSX-Portfolio.xlsx",
            )
            return
        if self.path.startswith("/api/export.pdf"):
            self.send_bytes(build_pdf(body), "application/pdf", "PSX-Portfolio.pdf")
            return
        self.send_error(404)

    def send_portfolio(self):
        if "latest" in self.path or _snapshot is None:
            refresh_in_background()
        if _snapshot:
            payload = dict(_snapshot)
            payload["refreshing"] = _refreshing
            self.send_json(payload)
            return
        status = market_status()
        self.send_json(
            {
                "refreshing": True,
                "rows": [],
                "updated": "",
                "market": status,
                "marketOpen": status.startswith("Market is open"),
                "stale": False,
                "error": _refresh_error,
            }
        )

    def send_bytes(self, raw, content_type, filename):
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.end_headers()
        self.wfile.write(raw)

    def send_json(self, payload, status=200):
        raw = json.dumps(payload).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)

    def log_message(self, fmt, *args):
        print("[%s] %s" % (stamp(), fmt % args))


def main():
    load_cache()
    host = "0.0.0.0"
    port = int(os.environ.get("PORT", "8765"))
    server = ThreadingHTTPServer((host, port), Handler)
    print("PSX Portfolio Tracker at http://%s:%s" % (host, port))
    server.serve_forever()


if __name__ == "__main__":
    main()
